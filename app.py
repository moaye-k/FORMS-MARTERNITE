import os
import csv
import datetime
from io import BytesIO
from flask import Flask, render_template, request, session, redirect, url_for

from reference_data import (
    INITIATEURS, INITIATEUR_SANS_OUTIL, PREMIERE_DEMANDE, MOMENT_DEMANDE,
    SUPPORT, SATISFACTION_ROWS, SATISFACTION_SCALE, CES, ERGONOMIE_ROWS,
    ERGONOMIE_SCALE, AIDE, CSAT, DELAI, CONFIANCE,
)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-moi-en-production")

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
RESPONSES_FILE = os.path.join(DATA_DIR, "reponses_maternite.csv")

CSV_HEADERS = (
    ["date_soumission", "initiateur", "utilise_outil_soi_meme",
     "premiere_demande", "moment_demande", "nps_note", "nps_raison", "support"]
    + [code for code, _ in SATISFACTION_ROWS]
    + ["ces"]
    + [code for code, _ in ERGONOMIE_ROWS]
    + ["aide", "csat", "delai", "confiance", "apprecie", "ameliorer"]
)


def save_response(data: dict):
    os.makedirs(DATA_DIR, exist_ok=True)
    file_exists = os.path.isfile(RESPONSES_FILE)
    with open(RESPONSES_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=CSV_HEADERS)
        if not file_exists:
            writer.writeheader()
        writer.writerow(data)


@app.route("/")
def home():
    session.clear()
    return render_template("home.html")


# ------------------------------------------------------------------
# Q1 : Qui a initié la demande ? -> détermine le parcours
# ------------------------------------------------------------------
@app.route("/q1", methods=["GET", "POST"])
def q1():
    if request.method == "POST":
        initiateur = request.form.get("initiateur")
        session["initiateur"] = initiateur
        session["utilise_outil_soi_meme"] = initiateur != INITIATEUR_SANS_OUTIL

        if session["utilise_outil_soi_meme"]:
            return redirect(url_for("q2"))
        else:
            # N'a pas utilise l'outil elle-meme -> saute la PARTIE 1
            return redirect(url_for("partie2"))

    return render_template("q1.html", initiateurs=INITIATEURS, step=1, total_steps=5)


# ------------------------------------------------------------------
# Q2 a Q5 : premiere demande, moment, NPS, raison de la note
# ------------------------------------------------------------------
@app.route("/q2", methods=["GET", "POST"])
def q2():
    if "initiateur" not in session:
        return redirect(url_for("q1"))

    if request.method == "POST":
        session["premiere_demande"] = request.form.get("premiere_demande")
        session["moment_demande"] = request.form.get("moment_demande")
        session["nps_note"] = request.form.get("nps_note")
        session["nps_raison"] = request.form.get("nps_raison", "")
        return redirect(url_for("partie1a"))

    return render_template(
        "q2.html",
        premiere_demande=PREMIERE_DEMANDE,
        moment_demande=MOMENT_DEMANDE,
        step=2, total_steps=5,
    )


# ------------------------------------------------------------------
# PARTIE 1a : Q6 support + Q7 matrice satisfaction par etape
# ------------------------------------------------------------------
@app.route("/partie1a", methods=["GET", "POST"])
def partie1a():
    if "initiateur" not in session:
        return redirect(url_for("q1"))

    if request.method == "POST":
        session["support"] = request.form.get("support")
        for code, _ in SATISFACTION_ROWS:
            session[code] = request.form.get(code, "")
        return redirect(url_for("partie1b"))

    return render_template(
        "partie1a.html",
        support=SUPPORT,
        rows=SATISFACTION_ROWS,
        scale=SATISFACTION_SCALE,
        step=3, total_steps=5,
    )


# ------------------------------------------------------------------
# PARTIE 1b : Q8 CES + Q9 matrice ergonomie
# ------------------------------------------------------------------
@app.route("/partie1b", methods=["GET", "POST"])
def partie1b():
    if "initiateur" not in session:
        return redirect(url_for("q1"))

    if request.method == "POST":
        session["ces"] = request.form.get("ces")
        for code, _ in ERGONOMIE_ROWS:
            session[code] = request.form.get(code, "")
        return redirect(url_for("partie1c"))

    return render_template(
        "partie1b.html",
        ces=CES,
        rows=ERGONOMIE_ROWS,
        scale=ERGONOMIE_SCALE,
        step=4, total_steps=5,
    )


# ------------------------------------------------------------------
# PARTIE 1c : Q10 aide recue
# ------------------------------------------------------------------
@app.route("/partie1c", methods=["GET", "POST"])
def partie1c():
    if "initiateur" not in session:
        return redirect(url_for("q1"))

    if request.method == "POST":
        session["aide"] = request.form.get("aide")
        return redirect(url_for("partie2"))

    return render_template("partie1c.html", aide=AIDE, step=4, total_steps=5)


# ------------------------------------------------------------------
# PARTIE 2 : Q11 a Q15 - remplie par TOUTES les repondantes
# ------------------------------------------------------------------
@app.route("/partie2", methods=["GET", "POST"])
def partie2():
    if "initiateur" not in session:
        return redirect(url_for("q1"))

    if request.method == "POST":
        response_data = {
            "date_soumission": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "initiateur": session.get("initiateur", ""),
            "utilise_outil_soi_meme": "Oui" if session.get("utilise_outil_soi_meme") else "Non",
            "premiere_demande": session.get("premiere_demande", ""),
            "moment_demande": session.get("moment_demande", ""),
            "nps_note": session.get("nps_note", ""),
            "nps_raison": session.get("nps_raison", ""),
            "support": session.get("support", ""),
            "ces": session.get("ces", ""),
            "aide": session.get("aide", ""),
            "csat": request.form.get("csat", ""),
            "delai": request.form.get("delai", ""),
            "confiance": request.form.get("confiance", ""),
            "apprecie": request.form.get("apprecie", ""),
            "ameliorer": request.form.get("ameliorer", ""),
        }
        for code, _ in SATISFACTION_ROWS:
            response_data[code] = session.get(code, "")
        for code, _ in ERGONOMIE_ROWS:
            response_data[code] = session.get(code, "")

        save_response(response_data)
        session.clear()
        return redirect(url_for("merci"))

    return render_template(
        "partie2.html",
        csat=CSAT, delai=DELAI, confiance=CONFIANCE,
        step=5, total_steps=5,
    )


@app.route("/merci")
def merci():
    return render_template("merci.html")


# ------------------------------------------------------------------
# Dashboard admin
# ------------------------------------------------------------------
@app.route("/admin/reponses")
def admin_reponses():
    admin_key = request.args.get("cle")
    if admin_key != os.environ.get("ADMIN_KEY", "maternite2026"):
        return "Accès refusé. Ajoutez ?cle=VOTRE_CLE à l'URL.", 403

    rows = []
    if os.path.isfile(RESPONSES_FILE):
        with open(RESPONSES_FILE, encoding="utf-8") as f:
            rows = list(csv.DictReader(f))

    total = len(rows)
    nps_scores = [int(r["nps_note"]) for r in rows if r.get("nps_note", "").isdigit()]
    nps_moyen = round(sum(nps_scores) / len(nps_scores), 1) if nps_scores else None

    # Calcul simplifie du NPS (%promoteurs - %detracteurs)
    promoteurs = len([s for s in nps_scores if s >= 9])
    detracteurs = len([s for s in nps_scores if s <= 6])
    nps_score = round((promoteurs - detracteurs) / len(nps_scores) * 100) if nps_scores else None

    csat_values = [r["csat"] for r in rows if r.get("csat")]
    csat_satisfaits = len([c for c in csat_values if c in ("Très satisfait(e)", "Satisfait(e)")])
    csat_pourcent = round(csat_satisfaits / len(csat_values) * 100) if csat_values else None

    return render_template(
        "admin.html",
        rows=rows, total=total, nps_moyen=nps_moyen,
        nps_score=nps_score, csat_pourcent=csat_pourcent,
    )


@app.route("/admin/reponses/export")
def admin_export():
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    admin_key = request.args.get("cle")
    if admin_key != os.environ.get("ADMIN_KEY", "maternite2026"):
        return "Accès refusé.", 403
    if not os.path.isfile(RESPONSES_FILE):
        return "Aucune réponse à exporter pour le moment.", 404

    with open(RESPONSES_FILE, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Réponses"
    worksheet.append(CSV_HEADERS)
    for row in rows:
        worksheet.append([row.get(header, "") for header in CSV_HEADERS])

    header_fill = PatternFill("solid", fgColor="0B2545")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="reponses_maternite.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
